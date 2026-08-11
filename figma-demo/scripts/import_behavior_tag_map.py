#!/usr/bin/env python3
"""Import 行为标签匹配表.xlsx → public/data/cz_behavior_tag_map.json"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path(__file__).resolve().parent / "data" / "行为标签匹配表.xlsx"
OUT = ROOT / "public" / "data" / "cz_behavior_tag_map.json"
OUT_SRC = ROOT / "src" / "lib" / "chronicle" / "data" / "cz_behavior_tag_map.json"

ASCII_IDS = {
    "早期入场": "early_entry",
    "后知后觉": "late_entry",
    "只碰主流": "bluechip_only",
    "爱冲新币": "chase_new",
    "链上老炮": "onchain_degen",
    "逆势加仓": "buy_dip",
    "追高买入": "fomo_buy",
    "全仓豪赌": "all_in",
    "分散保守": "diversified",
    "长期扛单": "long_hold",
    "割肉快": "cut_loss",
    "高抛低吸": "swing_trade",
    "拿不住": "cant_hold",
    "情绪驱动": "emotion_driven",
    "纪律型": "disciplined",
    "钝感扛压": "stoic",
    "爆过仓": "liquidated",
    "踩过雷": "rugged",
    "被套过": "underwater",
    "熬过熊": "survived_bear",
    "错过牛": "missed_bull",
    "离场又回归": "left_and_returned",
    "研究型": "researcher",
    "社交型": "social_follower",
    "早期押注": "early_bet",
    "快速转向": "fast_pivot",
    "长期主义": "long_termism",
    "追热点": "chase_narrative",
    "第一性思考": "first_principles",
    "以身作则": "lead_by_example",
    "善于放权": "delegate",
    "末位淘汰": "rank_and_yank",
    "重用产出": "output_first",
    "留人失败": "retention_fail",
    "简单直接": "keep_simple",
    "留退路": "exit_clause",
    "拒绝排他": "no_exclusivity",
    "被动 BD": "passive_bd",
    "踩过合作坑": "partner_burned",
    "用户优先": "user_first",
    "只做可规模化": "scalable_only",
    "过度打磨": "over_polish",
    "迭代快": "ship_fast",
    "兜底担责": "make_whole",
    "快速回应": "fast_response",
    "沉默吃亏": "silent_crisis",
    "扛过至暗": "near_zero",
    "被监管打过": "regulatory_hit",
    "冷静克制": "calm",
    "时间焦虑": "time_anxiety",
    "盛极转折": "peak_to_trough",
}


def slugify(label: str) -> str:
    if label in ASCII_IDS:
        return ASCII_IDS[label]
    return f"tag_{hashlib.sha1(label.encode('utf-8')).hexdigest()[:8]}"


def split_synonyms(cell: str) -> tuple[str, list[str]]:
    cell = (cell or "").strip()
    if not cell:
        return "", []
    parts = re.split(r"[，,]", cell, maxsplit=1)
    label = parts[0].strip()
    syns = [label]
    if len(parts) > 1:
        for piece in re.split(r"[、,/／\|]|以及", parts[1]):
            p = re.sub(r"^[·•\-\s]+", "", piece.strip())
            if p and p not in syns:
                syns.append(p)
    return label, syns


def main() -> None:
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx.is_file():
        print(f"missing xlsx: {xlsx}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["行为标签"]
    tags: list[dict] = []
    audience = "retail"
    dimension = ""
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        a, b, c = row[0], row[1], row[2]
        if a:
            audience = "founder" if "founder" in str(a).lower() else "retail"
        if b:
            dimension = str(b).strip()
        if not c:
            continue
        label, syns = split_synonyms(str(c))
        tags.append(
            {
                "id": slugify(label),
                "audience": [audience],
                "dimension": dimension,
                "label": label,
                "synonyms": syns,
            }
        )

    principles_by_tag: dict[str, list[dict]] = {}
    ws2 = wb["原则"]
    for i, row in enumerate(ws2.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        _a, b, c = row[0], row[1], row[2]
        if not b or not c:
            continue
        label = str(b).strip()
        tid = next((t["id"] for t in tags if t["label"] == label or label in t["synonyms"]), None)
        if not tid:
            tid = slugify(label)
        items = []
        for chunk in re.split(r"[／/]", str(c)):
            chunk = chunk.strip()
            if not chunk:
                continue
            m = re.match(r"^(.+?)（(.+?)）$", chunk)
            if m:
                items.append({"name": m.group(1).strip(), "note": m.group(2).strip()})
            else:
                items.append({"name": chunk, "note": None})
        principles_by_tag[tid] = items

    doc = {
        "version": 1,
        "source": xlsx.name,
        "tags": tags,
        "principlesByTag": principles_by_tag,
    }
    text = json.dumps(doc, ensure_ascii=False, indent=2) + "\n"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT_SRC.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(text, encoding="utf-8")
    OUT_SRC.write_text(text, encoding="utf-8")
    print(f"wrote {OUT} and {OUT_SRC} tags={len(tags)} mapped={len(principles_by_tag)}")


if __name__ == "__main__":
    main()
